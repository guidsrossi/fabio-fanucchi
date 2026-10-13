"""Extrai o mapão sem modificar os arquivos de origem."""
import argparse
import json
import re
import unicodedata
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('origem', type=Path)
args = parser.parse_args()
ws = openpyxl.load_workbook(args.origem, data_only=True, read_only=True).active
assert str(ws['B2'].value) == '2026'
assert 'Terceiro Bimestre' in str(ws['B7'].value)
descricao = unicodedata.normalize('NFD', str(ws['B6'].value)).encode('ascii', 'ignore').decode()
turma_match = re.search(r'([123])\s+SERIE\s+([A-E])', descricao)
assert turma_match, 'Turma não identificada'
turma = ''.join(turma_match.groups())
disciplinas = [(i, str(ws.cell(11, i).value).split('\n')[0]) for i in range(3, 63, 3)]
estudantes = []
for row in ws.iter_rows(min_row=13, values_only=True):
    if str(row[1]).strip().lower() != 'ativo':
        continue
    notas = []
    mencoes = []
    for col, disciplina in disciplinas:
        valor = row[col - 1]
        if valor is None or str(valor).strip() in ('', '-'):
            continue
        try:
            nota = float(str(valor).replace(',', '.'))
        except ValueError:
            mencoes.append({'disciplina': disciplina, 'mencao': str(valor).strip()})
            continue  # Menções de engajamento não são notas numéricas.
        assert 0 <= nota <= 10
        notas.append({'disciplina': disciplina, 'nota': nota, 'bimestre': 3})
    estudantes.append({'nome': str(row[0]).strip(), 'turma': turma, 'serie': f'{turma[0]}ª série',
                       'situacao': 'Ativo', 'frequencia': row[63],
                       'frequencia_anual': row[65], 'notas': notas, 'mencoes': mencoes,
                       'fonte': args.origem.name})
destino = Path(__file__).resolve().parents[1] / 'data/notas-terceiro-bimestre.json'
existentes = json.loads(destino.read_text(encoding='utf-8')) if destino.exists() else []
assert estudantes, 'Nenhum estudante ativo no mapão'
assert len({e['nome'] for e in estudantes}) == len(estudantes), 'Nomes duplicados'
todos = [e for e in existentes if e['turma'] != turma] + estudantes
todos.sort(key=lambda e: (e['turma'], e['nome']))
destino.write_text(json.dumps(todos, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'turma': turma, 'estudantes': len(estudantes), 'notas': sum(len(e['notas']) for e in estudantes)}))
